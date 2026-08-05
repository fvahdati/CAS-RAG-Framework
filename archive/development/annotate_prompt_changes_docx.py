#!/usr/bin/env python3
"""Annotate June 2026 Algorithm 1 prompt docx with line-level diffs vs current prompts."""

from __future__ import annotations

import difflib
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path(__file__).resolve().parents[1]
PROMPTS_DIR = ROOT / "resources/datasets/Experiments-data/video_02_education_tutoring_diner/prompts"
SOURCE_DOCX = PROMPTS_DIR / "Algorithm 1 latest run prompts_17June2026.docx"
OUTPUT_DOCX = PROMPTS_DIR / "Algorithm 1 latest run prompts_17June2026_CHANGES.docx"

# June docx section ranges (paragraph indices inclusive) -> current prompt file
SECTIONS = [
    {
        "title": "P1: Unknown Reference Detection (P1_unknown_reference_detection.txt)",
        "old_range": (0, 55),
        "current_file": "P1_unknown_reference_detection.txt",
    },
    {
        "title": "P2 (June) → Current P3: No-RAG Emotion with Ambiguous (P3_no_rag_emotion_classification.txt)",
        "old_range": (56, 113),
        "current_file": "P3_no_rag_emotion_classification.txt",
    },
    {
        "title": "P3 (June) → Current P7: No-RAG Emotion without Ambiguous (P7_no_rag_emotion_without_vagueness_check.txt)",
        "old_range": (114, 143),
        "current_file": "P7_no_rag_emotion_without_vagueness_check.txt",
    },
    {
        "title": "P4 (June) → Current P2: RAG Emotion Classification (P2_rag_emotion_classification.txt)",
        "old_range": (144, 198),
        "current_file": "P2_rag_emotion_classification.txt",
    },
    {
        "title": "P5 (June) → Current P6: Semantic Ambiguity-Aware Emotion (P6_semantic_ambiguity_aware_emotion.txt)",
        "old_range": (199, 274),
        "current_file": "P6_semantic_ambiguity_aware_emotion.txt",
    },
    {
        "title": "P6 (June) → Current P5: Semantic Similarity Explicit/Ambiguous (P5_semantic_similarity_explicit_vague.txt)",
        "old_range": (275, 355),
        "current_file": "P5_semantic_similarity_explicit_vague.txt",
    },
]

SKIP_OLD_PREFIXES = (
    "Algorithm 1 latest run prompts:",
    "P1: Unknown Reference Detection with LLM",
    "P2. No-RAG Emotion Classification Prompt",
    "P3. No-RAG Emotion Classification Prompt",
    "P4. RAG-Based Emotion Classification Prompt",
    "P5.  Semantic Ambiguity-Aware Emotion Classification Prompt",
    "P6. Semantic Similarity Computation",
    "Semantic Similarity Computation Prompt",
    "No-RAG Emotion Classification Prompt (without checking vagueness)",
)


def _normalize(line: str) -> str:
    return " ".join(line.split())


def extract_old_lines(doc: Document, start: int, end: int) -> list[str]:
    lines: list[str] = []
    for p in doc.paragraphs[start : end + 1]:
        text = p.text.strip()
        if not text:
            continue
        if any(text.startswith(prefix) for prefix in SKIP_OLD_PREFIXES):
            continue
        # Preserve multi-line paragraphs split in docx
        for part in text.split("\n"):
            part = part.strip()
            if part:
                lines.append(part)
    return lines


def load_current_lines(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    return [ln.rstrip() for ln in text.splitlines() if ln.strip()]


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_diff_line(doc: Document, prefix: str, text: str, *, red: bool = False, strike: bool = False) -> None:
    para = doc.add_paragraph()
    run = para.add_run(f"{prefix}{text}")
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    if red:
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)  # dark red
        run.bold = True
    if strike:
        run.font.strike = True
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def build_diff_section(doc: Document, old_lines: list[str], new_lines: list[str]) -> tuple[int, int, int]:
    """Return counts: removed, added, unchanged."""
    matcher = difflib.SequenceMatcher(None, old_lines, new_lines, autojunk=False)
    removed = added = unchanged = 0

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for line in old_lines[i1:i2]:
                add_diff_line(doc, "  ", line)
                unchanged += 1
        elif tag == "replace":
            for line in old_lines[i1:i2]:
                add_diff_line(doc, "- ", line, strike=True)
                removed += 1
            for line in new_lines[j1:j2]:
                add_diff_line(doc, "+ ", line, red=True)
                added += 1
        elif tag == "delete":
            for line in old_lines[i1:i2]:
                add_diff_line(doc, "- ", line, strike=True)
                removed += 1
        elif tag == "insert":
            for line in new_lines[j1:j2]:
                add_diff_line(doc, "+ ", line, red=True)
                added += 1
    return removed, added, unchanged


def main() -> None:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)

    source = Document(SOURCE_DOCX)
    out = Document()

    out.add_heading("Algorithm 1 Prompt Changes — June 17 2026 → Current", level=0)
    intro = out.add_paragraph()
    intro.add_run(
        "Legend: lines prefixed with "
    )
    minus = intro.add_run("- ")
    minus.font.strike = True
    minus.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    intro.add_run("were removed; lines prefixed with ")
    plus = intro.add_run("+ ")
    plus.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    plus.bold = True
    intro.add_run("are new or changed (shown in red). Unchanged lines have no prefix.")

    summary_rows: list[tuple[str, int, int, int]] = []

    for section in SECTIONS:
        add_heading(out, section["title"], level=1)
        old_lines = extract_old_lines(source, *section["old_range"])
        current_path = PROMPTS_DIR / section["current_file"]
        new_lines = load_current_lines(current_path)

        note = out.add_paragraph()
        note.add_run(f"Compared {len(old_lines)} old lines vs {len(new_lines)} current lines.\n")

        removed, added, unchanged = build_diff_section(out, old_lines, new_lines)
        summary_rows.append((section["title"], removed, added, unchanged))
        out.add_paragraph()

    add_heading(out, "Summary", level=1)
    table = out.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Prompt section"
    hdr[1].text = "Removed (-)"
    hdr[2].text = "Added/changed (+)"
    hdr[3].text = "Unchanged"
    for title, rem, add, same in summary_rows:
        row = table.add_row().cells
        row[0].text = title.split("(")[0].strip()
        row[1].text = str(rem)
        row[2].text = str(add)
        row[3].text = str(same)

    out.save(OUTPUT_DOCX)
    print(f"Wrote: {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
